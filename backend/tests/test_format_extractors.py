"""Tests for Module 1.4 — format-specific extractors (Text, DOCX, PPTX, EPUB, URL, Kreuzberg)."""

from __future__ import annotations

import asyncio
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from app.services.extraction.base import (
    ExtractionError,
    ExtractorDependencyError,
)
from app.services.extraction.docx import DocxExtractor, _paragraph_heading_level
from app.services.extraction.models import SourceInput
from app.services.extraction.text import PlainTextExtractor

# ============================================================================
# PlainTextExtractor
# ============================================================================


class TestPlainTextExtractor:
    def test_can_handle_txt_extension(self) -> None:
        e = PlainTextExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", filename="readme.txt")) is True

    def test_can_handle_md_extension(self) -> None:
        e = PlainTextExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", filename="notes.md")) is True

    def test_can_handle_text_source_type(self) -> None:
        e = PlainTextExtractor()
        assert e.can_handle(SourceInput(source_type="text")) is True

    def test_can_handle_raw_text(self) -> None:
        e = PlainTextExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", raw_text="hello")) is True

    def test_rejects_pdf(self) -> None:
        e = PlainTextExtractor()
        assert e.can_handle(SourceInput(source_type="pdf", filename="doc.pdf")) is False

    def test_extract_raw_text(self) -> None:
        e = PlainTextExtractor()
        source = SourceInput(source_type="text", raw_text="Hello, world!")
        doc = asyncio.run(e.extract(source))
        assert doc.full_text == "Hello, world!"
        assert doc.markdown is None
        assert doc.extraction_engine == "plain_text"

    def test_extract_markdown_file(self, tmp_path: Path) -> None:
        md_file = tmp_path / "notes.md"
        md_file.write_text("# Title\n\nBody.", encoding="utf-8")
        e = PlainTextExtractor()
        source = SourceInput(
            source_type="text",
            file_path=str(md_file),
            filename="notes.md",
        )
        doc = asyncio.run(e.extract(source))
        assert doc.markdown is not None
        assert "# Title" in doc.markdown

    def test_extract_txt_file(self, tmp_path: Path) -> None:
        txt_file = tmp_path / "notes.txt"
        txt_file.write_text("Plain text file.", encoding="utf-8")
        e = PlainTextExtractor()
        source = SourceInput(source_type="text", file_path=str(txt_file), filename="notes.txt")
        doc = asyncio.run(e.extract(source))
        assert doc.markdown is None
        assert doc.full_text == "Plain text file."

    def test_extract_empty_text_raises(self) -> None:
        e = PlainTextExtractor()
        with pytest.raises(ExtractionError):
            asyncio.run(e.extract(SourceInput(source_type="text", raw_text="   ")))

    def test_pasted_text_display_name(self) -> None:
        e = PlainTextExtractor()
        doc = asyncio.run(e.extract(SourceInput(source_type="text", raw_text="x")))
        assert doc.title == "Pasted Text"


# ============================================================================
# paragraph_heading_level (DOCX helper)
# ============================================================================


class TestParagraphHeadingLevel:
    def test_heading_1(self) -> None:
        assert _paragraph_heading_level("Heading 1") == 1

    def test_heading_3(self) -> None:
        assert _paragraph_heading_level("heading 3") == 3

    def test_normal_style(self) -> None:
        assert _paragraph_heading_level("Normal") is None

    def test_empty_string(self) -> None:
        assert _paragraph_heading_level("") is None

    def test_garbage_style(self) -> None:
        assert _paragraph_heading_level("Heading X") is None


# ============================================================================
# DOCX (mocked)
# ============================================================================


class TestDocxExtractor:
    def test_can_handle_docx_extension(self) -> None:
        e = DocxExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", filename="report.docx")) is True

    def test_can_handle_docx_source_type(self) -> None:
        e = DocxExtractor()
        assert e.can_handle(SourceInput(source_type="docx")) is True

    def test_rejects_pptx(self) -> None:
        e = DocxExtractor()
        assert e.can_handle(SourceInput(source_type="pptx")) is False

    def test_missing_file_path_raises(self) -> None:
        e = DocxExtractor()
        with pytest.raises(ExtractionError):
            asyncio.run(e.extract(SourceInput(source_type="docx")))

    def test_missing_dependency_raises(self) -> None:
        e = DocxExtractor()
        with (
            patch.dict("sys.modules", {"docx": None}),
            pytest.raises(ExtractorDependencyError, match="python-docx"),
        ):
            asyncio.run(e.extract(SourceInput(
                source_type="docx", file_path="/tmp/test.docx",
            )))

    def test_extract_with_headings(self, tmp_path: Path) -> None:
        """Use python-docx to create a real DOCX and extract it."""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_heading("Chapter 1", level=1)
        doc.add_paragraph("Body paragraph one.")
        doc.add_heading("Section 1.1", level=2)
        doc.add_paragraph("Body paragraph two.")
        doc.save(str(docx_path))

        e = DocxExtractor()
        source = SourceInput(
            source_type="docx",
            file_path=str(docx_path),
            filename="test.docx",
        )
        result = asyncio.run(e.extract(source))
        assert result.full_text
        assert result.markdown is not None
        assert "# Chapter 1" in result.markdown
        assert "## Section 1.1" in result.markdown
        assert "Body paragraph one" in result.markdown
        assert result.metadata["source_type"] == "docx"

    def test_empty_docx_raises(self, tmp_path: Path) -> None:
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "empty.docx"
        doc = Document()
        # No paragraphs — entirely empty
        doc.save(str(docx_path))

        e = DocxExtractor()
        with pytest.raises(ExtractionError, match="No text"):
            asyncio.run(e.extract(SourceInput(
                source_type="docx", file_path=str(docx_path),
            )))


# ============================================================================
# PPTX (mocked and real)
# ============================================================================


class TestPptxExtractor:
    def test_can_handle_pptx_extension(self) -> None:
        from app.services.extraction.pptx import PptxExtractor
        e = PptxExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", filename="deck.pptx")) is True

    def test_can_handle_pptx_source_type(self) -> None:
        from app.services.extraction.pptx import PptxExtractor
        e = PptxExtractor()
        assert e.can_handle(SourceInput(source_type="pptx")) is True

    def test_missing_file_path_raises(self) -> None:
        from app.services.extraction.pptx import PptxExtractor
        e = PptxExtractor()
        with pytest.raises(ExtractionError):
            asyncio.run(e.extract(SourceInput(source_type="pptx")))

    def test_missing_dependency_raises(self) -> None:
        from app.services.extraction.pptx import PptxExtractor
        e = PptxExtractor()
        with (
            patch.dict("sys.modules", {"pptx": None}),
            pytest.raises(ExtractorDependencyError, match="python-pptx"),
        ):
            asyncio.run(e.extract(SourceInput(
                source_type="pptx", file_path="/tmp/test.pptx",
            )))

    def test_extract_two_slides(self, tmp_path: Path) -> None:
        try:
            from pptx import Presentation  # noqa: F401
        except ImportError:
            pytest.skip("python-pptx not installed")

        pptx_path = tmp_path / "test.pptx"
        prs = Presentation()
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        if slide1.shapes.title:
            slide1.shapes.title.text = "Slide 1 Title"
        body = slide1.placeholders[1]
        body.text = "Slide 1 body"

        slide2 = prs.slides.add_slide(prs.slide_layouts[0])
        if slide2.shapes.title:
            slide2.shapes.title.text = "Slide 2 Title"
        body2 = slide2.placeholders[1]
        body2.text = "Slide 2 body"

        prs.save(str(pptx_path))

        from app.services.extraction.pptx import PptxExtractor
        e = PptxExtractor()
        source = SourceInput(
            source_type="pptx",
            file_path=str(pptx_path),
            filename="test.pptx",
        )
        result = asyncio.run(e.extract(source))
        assert result.slides is not None
        assert len(result.slides) >= 2
        assert result.metadata["source_type"] == "pptx"
        assert result.extraction_engine == "python-pptx"


# ============================================================================
# EPUB (mocked)
# ============================================================================


class TestEpubExtractor:
    def test_can_handle_epub_extension(self) -> None:
        from app.services.extraction.epub import EpubExtractor
        e = EpubExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", filename="book.epub")) is True

    def test_can_handle_epub_source_type(self) -> None:
        from app.services.extraction.epub import EpubExtractor
        e = EpubExtractor()
        assert e.can_handle(SourceInput(source_type="epub")) is True

    def test_missing_file_path_raises(self) -> None:
        from app.services.extraction.epub import EpubExtractor
        e = EpubExtractor()
        with pytest.raises(ExtractionError):
            asyncio.run(e.extract(SourceInput(source_type="epub")))

    def test_missing_dependency_raises(self) -> None:
        from app.services.extraction.epub import EpubExtractor
        e = EpubExtractor()
        with (
            patch.dict("sys.modules", {"ebooklib": None}),
            pytest.raises(ExtractorDependencyError, match="ebooklib"),
        ):
            asyncio.run(e.extract(SourceInput(
                source_type="epub", file_path="/tmp/test.epub",
            )))


# ============================================================================
# EPUB TOC href injection helpers (unit tests)
# ============================================================================


class TestNormalizeHrefPath:
    def test_lowercase(self) -> None:
        from app.services.extraction.epub import _normalize_href_path
        assert _normalize_href_path("Text/Ch01.XHTML") == "text/ch01.xhtml"

    def test_backslash_normalization(self) -> None:
        from app.services.extraction.epub import _normalize_href_path
        assert _normalize_href_path("Text\\ch01.xhtml") == "text/ch01.xhtml"

    def test_strips_leading_trailing_slashes(self) -> None:
        from app.services.extraction.epub import _normalize_href_path
        assert _normalize_href_path("/Text/ch01.xhtml/") == "text/ch01.xhtml"

    def test_whitespace_trim(self) -> None:
        from app.services.extraction.epub import _normalize_href_path
        assert _normalize_href_path("  Text/ch01.xhtml  ") == "text/ch01.xhtml"


class TestFlattenTocForInjection:
    def test_empty_toc(self) -> None:
        from app.services.extraction.epub import _flatten_toc_for_injection
        assert _flatten_toc_for_injection([]) == []

    def test_single_link(self) -> None:
        from app.services.extraction.epub import _flatten_toc_for_injection
        from ebooklib import epub
        toc = [epub.Link("Text/ch01.xhtml", "CHAPTER I")]
        result = _flatten_toc_for_injection(toc)
        assert len(result) == 1
        assert result[0]["title"] == "CHAPTER I"
        assert result[0]["href_path"] == "Text/ch01.xhtml"
        assert result[0]["fragment"] == ""

    def test_link_with_fragment(self) -> None:
        from app.services.extraction.epub import _flatten_toc_for_injection
        from ebooklib import epub
        toc = [epub.Link("Text/book.xhtml#ch01", "Chapter 1")]
        result = _flatten_toc_for_injection(toc)
        assert len(result) == 1
        assert result[0]["title"] == "Chapter 1"
        assert result[0]["href_path"] == "Text/book.xhtml"
        assert result[0]["fragment"] == "ch01"

    def test_tuple_toc(self) -> None:
        from app.services.extraction.epub import _flatten_toc_for_injection
        from ebooklib import epub
        toc = [(epub.Link("a.xhtml", "Parent"), [epub.Link("b.xhtml", "Child")])]
        result = _flatten_toc_for_injection(toc)
        assert len(result) == 2
        assert result[0]["title"] == "Parent"
        assert result[1]["title"] == "Child"

    def test_section_without_href_is_skipped(self) -> None:
        from app.services.extraction.epub import _flatten_toc_for_injection
        from ebooklib import epub
        toc = [epub.Section("Part I", href="")]
        result = _flatten_toc_for_injection(toc)
        assert len(result) == 0

    def test_section_with_href_is_included(self) -> None:
        from app.services.extraction.epub import _flatten_toc_for_injection
        from ebooklib import epub
        toc = [epub.Section("Part I", href="part1.xhtml")]
        result = _flatten_toc_for_injection(toc)
        assert len(result) == 1
        assert result[0]["title"] == "Part I"

    def test_nested_list_toc(self) -> None:
        from app.services.extraction.epub import _flatten_toc_for_injection
        from ebooklib import epub
        toc = [
            epub.Link("ch01.xhtml", "Chapter 1"),
            (epub.Section("Part II", href="part2.xhtml"), [
                epub.Link("ch02.xhtml", "Chapter 2"),
                epub.Link("ch03.xhtml", "Chapter 3"),
            ]),
        ]
        result = _flatten_toc_for_injection(toc)
        assert len(result) == 4
        titles = [r["title"] for r in result]
        assert titles == ["Chapter 1", "Part II", "Chapter 2", "Chapter 3"]


class TestMatchItemToTocTitle:
    def test_suffix_match_ncx_style(self) -> None:
        from app.services.extraction.epub import _match_item_to_toc_title
        entries = [{"title": "CHAPTER I", "href_path": "Text/ch01.xhtml", "fragment": ""}]
        assert _match_item_to_toc_title("OEBPS/Text/ch01.xhtml", entries) == "CHAPTER I"

    def test_suffix_match_nav_style(self) -> None:
        from app.services.extraction.epub import _match_item_to_toc_title
        entries = [{"title": "CHAPTER I", "href_path": "OEBPS/Text/ch01.xhtml", "fragment": ""}]
        assert _match_item_to_toc_title("OEBPS/Text/ch01.xhtml", entries) == "CHAPTER I"

    def test_no_match(self) -> None:
        from app.services.extraction.epub import _match_item_to_toc_title
        entries = [{"title": "CHAPTER I", "href_path": "Text/other.xhtml", "fragment": ""}]
        assert _match_item_to_toc_title("OEBPS/Text/ch01.xhtml", entries) is None

    def test_first_match_wins(self) -> None:
        from app.services.extraction.epub import _match_item_to_toc_title
        entries = [
            {"title": "Title Page", "href_path": "Text/content.xhtml", "fragment": ""},
            {"title": "CHAPTER I", "href_path": "Text/content.xhtml", "fragment": "ch01"},
        ]
        assert _match_item_to_toc_title("OEBPS/Text/content.xhtml", entries) == "Title Page"

    def test_empty_entries_list(self) -> None:
        from app.services.extraction.epub import _match_item_to_toc_title
        assert _match_item_to_toc_title("OEBPS/Text/ch01.xhtml", []) is None


class TestHtmlHasHeadingWithTitle:
    def test_h1_match(self) -> None:
        from app.services.extraction.epub import _html_has_heading_with_title
        html = "<html><body><h1>Chapter I</h1></body></html>"
        assert _html_has_heading_with_title(html, "Chapter I") is True

    def test_case_insensitive(self) -> None:
        from app.services.extraction.epub import _html_has_heading_with_title
        html = "<h1>chapter i</h1>"
        assert _html_has_heading_with_title(html, "CHAPTER I") is True

    def test_h2_match(self) -> None:
        from app.services.extraction.epub import _html_has_heading_with_title
        html = "<h2>Section 1</h2>"
        assert _html_has_heading_with_title(html, "Section 1") is True

    def test_no_heading_match(self) -> None:
        from app.services.extraction.epub import _html_has_heading_with_title
        html = "<p>Chapter I</p>"
        assert _html_has_heading_with_title(html, "Chapter I") is False

    def test_partial_no_match(self) -> None:
        from app.services.extraction.epub import _html_has_heading_with_title
        html = "<h1>Chapter I: Intro</h1>"
        assert _html_has_heading_with_title(html, "Chapter I") is False

    def test_nested_tags_in_heading(self) -> None:
        from app.services.extraction.epub import _html_has_heading_with_title
        html = "<h1>Chapter <em>I</em></h1>"
        assert _html_has_heading_with_title(html, "Chapter I") is True


class TestInjectUnmatchedTitles:
    def test_simple_injection(self) -> None:
        from app.services.extraction.epub import _inject_unmatched_titles
        entries = [{"title": "CHAPTER II", "href_path": "ch02.xhtml", "fragment": ""}]
        md = "Some text.\n\nCHAPTER II\n\nMore text."
        result = _inject_unmatched_titles(md, entries, set())
        assert "# CHAPTER II" in result

    def test_case_insensitive_match(self) -> None:
        from app.services.extraction.epub import _inject_unmatched_titles
        entries = [{"title": "CHAPTER II", "href_path": "ch02.xhtml", "fragment": ""}]
        md = "chapter ii\n\nBody."
        result = _inject_unmatched_titles(md, entries, set())
        assert "# chapter ii" in result

    def test_skip_already_heading(self) -> None:
        from app.services.extraction.epub import _inject_unmatched_titles
        entries = [{"title": "CHAPTER II", "href_path": "ch02.xhtml", "fragment": ""}]
        md = "# CHAPTER II\n\nBody."
        result = _inject_unmatched_titles(md, entries, set())
        # Should remain unchanged — already a heading
        assert result == md

    def test_only_first_occurrence(self) -> None:
        from app.services.extraction.epub import _inject_unmatched_titles
        entries = [{"title": "CHAPTER II", "href_path": "ch02.xhtml", "fragment": ""}]
        md = "CHAPTER II\n\nBody.\n\nCHAPTER II"
        result = _inject_unmatched_titles(md, entries, set())
        assert result.count("# CHAPTER II") == 1

    def test_skip_injected_via_href(self) -> None:
        from app.services.extraction.epub import _inject_unmatched_titles
        entries = [{"title": "CHAPTER I", "href_path": "ch01.xhtml", "fragment": ""}]
        md = "CHAPTER I\n\nBody."
        result = _inject_unmatched_titles(md, entries, {"chapter i"})
        assert result == md  # unchanged

    def test_whitespace_handling(self) -> None:
        from app.services.extraction.epub import _inject_unmatched_titles
        entries = [{"title": "CHAPTER II", "href_path": "ch02.xhtml", "fragment": ""}]
        md = "  CHAPTER II  \n\nBody."
        result = _inject_unmatched_titles(md, entries, set())
        assert "# CHAPTER II" in result

    def test_empty_markdown(self) -> None:
        from app.services.extraction.epub import _inject_unmatched_titles
        entries = [{"title": "CHAPTER I", "href_path": "ch01.xhtml", "fragment": ""}]
        assert _inject_unmatched_titles("", entries, set()) == ""


# ============================================================================
# EPUB integration tests (mocked ebooklib with TOC injection)
# ============================================================================


class TestEpubExtractorWithTocInjection:
    def test_multi_file_epub_css_titles(self) -> None:
        """Layer 1: multi-file EPUB where each chapter file is matched by href."""
        from app.services.extraction.epub import EpubExtractor
        from ebooklib import epub

        mock_book = MagicMock()
        mock_book.get_metadata.return_value = []
        mock_book.toc = [
            epub.Link("Text/ch01.xhtml", "CHAPTER I"),
            epub.Link("Text/ch02.xhtml", "CHAPTER II"),
            epub.Link("Text/ch03.xhtml", "CHAPTER III"),
        ]

        def make_item(name: str, html: str) -> MagicMock:
            item = MagicMock()
            item.get_name.return_value = name
            item.get_content.return_value = html.encode("utf-8")
            return item

        items = [
            make_item("OEBPS/Text/ch01.xhtml", "<p class=\"chapter-title\">CHAPTER I</p><p>Text one.</p>"),
            make_item("OEBPS/Text/ch02.xhtml", "<p class=\"chapter-title\">CHAPTER II</p><p>Text two.</p>"),
            make_item("OEBPS/Text/ch03.xhtml", "<p class=\"chapter-title\">CHAPTER III</p><p>Text three.</p>"),
        ]
        mock_book.get_items_of_type.return_value = items

        with patch("ebooklib.epub.read_epub", return_value=mock_book):
            e = EpubExtractor()
            doc = asyncio.run(e.extract(SourceInput(
                source_type="epub", file_path="/tmp/test.epub",
            )))

        assert doc.markdown is not None
        assert "# CHAPTER I" in doc.markdown
        assert "# CHAPTER II" in doc.markdown
        assert "# CHAPTER III" in doc.markdown

    def test_single_file_epub_with_fragments(self) -> None:
        """Layer 1 + 2: single-file EPUB — first title from href, rest from text scan."""
        from app.services.extraction.epub import EpubExtractor
        from ebooklib import epub

        mock_book = MagicMock()
        mock_book.get_metadata.return_value = []
        mock_book.toc = [
            epub.Link("Text/book.xhtml#ch01", "CHAPTER I"),
            epub.Link("Text/book.xhtml#ch02", "CHAPTER II"),
        ]

        html = (
            "<p class=\"chapter-title\">CHAPTER I</p>\n"
            "<p>First chapter body text here.</p>\n"
            "<p>CHAPTER II</p>\n"
            "<p>Second chapter body text here.</p>"
        )
        item = MagicMock()
        item.get_name.return_value = "OEBPS/Text/book.xhtml"
        item.get_content.return_value = html.encode("utf-8")
        mock_book.get_items_of_type.return_value = [item]

        with patch("ebooklib.epub.read_epub", return_value=mock_book):
            e = EpubExtractor()
            doc = asyncio.run(e.extract(SourceInput(
                source_type="epub", file_path="/tmp/test.epub",
            )))

        assert doc.markdown is not None
        # Layer 1: href-matched title injected via <h1> prepend
        assert "# CHAPTER I" in doc.markdown
        # Layer 2: text-scanning found "CHAPTER II" line
        assert "# CHAPTER II" in doc.markdown

    def test_think_python_no_regression(self) -> None:
        """EPUBs with real <h1> tags are not double-injected."""
        from app.services.extraction.epub import EpubExtractor
        from ebooklib import epub

        mock_book = MagicMock()
        mock_book.get_metadata.return_value = []
        mock_book.toc = [
            epub.Link("Text/ch01.xhtml", "Chapter 1. The Way of the Program"),
        ]

        html = "<h1>Chapter 1. The Way of the Program</h1>\n<p>Python is fun.</p>"
        item = MagicMock()
        item.get_name.return_value = "OEBPS/Text/ch01.xhtml"
        item.get_content.return_value = html.encode("utf-8")
        mock_book.get_items_of_type.return_value = [item]

        with patch("ebooklib.epub.read_epub", return_value=mock_book):
            e = EpubExtractor()
            doc = asyncio.run(e.extract(SourceInput(
                source_type="epub", file_path="/tmp/test.epub",
            )))

        assert doc.markdown is not None
        # The <h1> is already converted to "# Chapter 1. The Way of the Program"
        # No double-injection — should appear exactly once
        heading_line = "# Chapter 1. The Way of the Program"
        assert doc.markdown.count(heading_line) == 1

    def test_empty_toc_no_injection(self) -> None:
        """EPUBs with empty TOC produce unchanged output."""
        from app.services.extraction.epub import EpubExtractor

        mock_book = MagicMock()
        mock_book.get_metadata.return_value = []
        mock_book.toc = []

        html = "<p class=\"chapter-title\">CHAPTER I</p><p>Body.</p>"
        item = MagicMock()
        item.get_name.return_value = "OEBPS/Text/ch01.xhtml"
        item.get_content.return_value = html.encode("utf-8")
        mock_book.get_items_of_type.return_value = [item]

        with patch("ebooklib.epub.read_epub", return_value=mock_book):
            e = EpubExtractor()
            doc = asyncio.run(e.extract(SourceInput(
                source_type="epub", file_path="/tmp/test.epub",
            )))

        assert doc.markdown is not None
        # No # headings injected since TOC is empty
        assert "# CHAPTER I" not in doc.markdown


# ============================================================================
# URL (mocked trafilatura)
# ============================================================================


class TestUrlExtractor:
    def test_can_handle_url_source_type(self) -> None:
        from app.services.extraction.url import UrlExtractor
        e = UrlExtractor()
        assert e.can_handle(SourceInput(source_type="url", url="https://example.com")) is True

    def test_can_handle_url_in_other_type(self) -> None:
        from app.services.extraction.url import UrlExtractor
        e = UrlExtractor()
        assert e.can_handle(SourceInput(
            source_type="unknown",
            url="https://example.com/article",
        )) is True

    def test_rejects_non_http(self) -> None:
        from app.services.extraction.url import UrlExtractor
        e = UrlExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", url="ftp://files.com")) is False

    def test_missing_url_raises(self) -> None:
        from app.services.extraction.url import UrlExtractor
        e = UrlExtractor()
        with pytest.raises(ExtractionError):
            asyncio.run(e.extract(SourceInput(source_type="url")))

    def test_invalid_url_protocol_raises(self) -> None:
        from app.services.extraction.url import UrlExtractor
        e = UrlExtractor()
        with pytest.raises(ExtractionError, match="http"):
            asyncio.run(e.extract(SourceInput(source_type="url", url="ftp://files.com")))

    def test_mock_extraction(self) -> None:
        from app.services.extraction.url import UrlExtractor

        with (
            patch("trafilatura.fetch_url", return_value="<html><body><p>Article text</p></body></html>"),
            patch("trafilatura.extract", return_value="Article text"),
            patch("trafilatura.metadata.extract_metadata", return_value=None),
        ):
            e = UrlExtractor()
            doc = asyncio.run(e.extract(SourceInput(
                source_type="url",
                url="https://example.com/article",
            )))
            assert doc.full_text == "Article text"
            assert doc.metadata["url"] == "https://example.com/article"

    def test_failed_fetch_raises(self) -> None:
        from app.services.extraction.url import UrlExtractor

        with patch("trafilatura.fetch_url", return_value=None):
            e = UrlExtractor()
            with pytest.raises(ExtractionError, match="fetch"):
                asyncio.run(e.extract(SourceInput(
                    source_type="url", url="https://example.com",
                )))


# ============================================================================
# Kreuzberg fallback
# ============================================================================


class TestKreuzbergFallback:
    def test_accepts_file_backed_sources(self) -> None:
        from app.services.extraction.kreuzberg_fallback import KreuzbergFallbackExtractor
        e = KreuzbergFallbackExtractor()
        assert e.can_handle(SourceInput(source_type="unknown", file_path="/tmp/file.xyz")) is True

    def test_rejects_no_file_path(self) -> None:
        from app.services.extraction.kreuzberg_fallback import KreuzbergFallbackExtractor
        e = KreuzbergFallbackExtractor()
        assert e.can_handle(SourceInput(source_type="text", raw_text="x")) is False

    def test_missing_file_path_raises(self) -> None:
        from app.services.extraction.kreuzberg_fallback import KreuzbergFallbackExtractor
        e = KreuzbergFallbackExtractor()
        with pytest.raises(ExtractionError):
            asyncio.run(e.extract(SourceInput(source_type="unknown")))

    def test_missing_dependency_raises(self) -> None:
        from app.services.extraction.kreuzberg_fallback import KreuzbergFallbackExtractor
        e = KreuzbergFallbackExtractor()
        with (
            patch.dict("sys.modules", {"kreuzberg": None}),
            pytest.raises(ExtractorDependencyError, match="kreuzberg"),
        ):
            asyncio.run(e.extract(SourceInput(
                source_type="unknown", file_path="/tmp/test.xyz",
            )))

    def test_fallback_warning_included(self) -> None:
        fake_result = MagicMock()
        fake_result.content = "Extracted content"
        mock_kb = MagicMock()
        mock_kb.extract_file = AsyncMock(return_value=fake_result)

        from app.services.extraction.kreuzberg_fallback import KreuzbergFallbackExtractor

        with patch.dict("sys.modules", {"kreuzberg": mock_kb}):
            e = KreuzbergFallbackExtractor()
            doc = asyncio.run(e.extract(SourceInput(
                source_type="unknown",
                file_path="/tmp/test.xyz",
                filename="test.xyz",
            )))
            assert doc.full_text == "Extracted content"
            assert len(doc.warnings) >= 1
            assert any("fallback" in w.lower() for w in doc.warnings)
            assert doc.extraction_engine == "kreuzberg"


# ============================================================================
# Dependency error message check
# ============================================================================


class TestDependencyErrors:
    def test_error_messages_mention_requirements_file(self) -> None:
        err = ExtractorDependencyError(
            "python-docx is required. pip install -r backend/requirements-extractors.txt"
        )
        assert "requirements-extractors.txt" in str(err)
        assert "python-docx" in str(err)


# ============================================================================
# Module 1.10 — Dependency file check
# ============================================================================


class TestRequirementsExtractorsFile:
    def test_file_exists(self) -> None:
        import os
        req_path = os.path.join(
            os.path.dirname(__file__), "..", "requirements-extractors.txt"
        )
        assert os.path.isfile(req_path), f"Missing: {req_path}"

    def test_contains_required_packages(self) -> None:
        import os
        req_path = os.path.join(
            os.path.dirname(__file__), "..", "requirements-extractors.txt"
        )
        content = open(req_path).read()
        assert "ebooklib" in content
        assert "python-pptx" in content
        assert "python-docx" in content
        assert "kreuzberg" in content
