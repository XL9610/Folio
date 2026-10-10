"""DOCX extractor — preserves Word heading hierarchy as Markdown."""

from __future__ import annotations

from pathlib import Path

from .base import ContentExtractor, ExtractionError, ExtractorDependencyError
from .models import ExtractedDocument, SourceInput

_INSTALL_HINT = "pip install -r backend/requirements-extractors.txt"


class DocxExtractor(ContentExtractor):
    """Extract DOCX with heading styles preserved as ATX markdown headings."""

    engine_name = "python-docx"

    # ------------------------------------------------------------------
    # ContentExtractor interface
    # ------------------------------------------------------------------

    def can_handle(self, source: SourceInput) -> bool:
        ext = Path(source.filename or source.file_path or "").suffix.lower()
        return source.source_type == "docx" or ext == ".docx"

    async def extract(self, source: SourceInput) -> ExtractedDocument:
        file_path = source.file_path
        if not file_path:
            raise ExtractionError("DOCX extraction requires a file_path")

        # Optional dependency guard
        try:
            from docx import Document
        except ImportError:
            raise ExtractorDependencyError(
                f"python-docx is required for DOCX extraction. Install with: {_INSTALL_HINT}"
            ) from None

        doc = Document(file_path)
        markdown_lines: list[str] = []
        text_lines: list[str] = []
        heading_count = 0
        para_count = 0

        for para in doc.paragraphs:
            level = _paragraph_heading_level(para.style.name if para.style else "")
            content = para.text.strip()
            if not content:
                continue

            para_count += 1
            if level is not None:
                heading_count += 1
                markdown_lines.append(f"{'#' * level} {content}")
            else:
                markdown_lines.append(content)

            text_lines.append(content)

        full_text = "\n".join(text_lines)
        if not full_text.strip():
            raise ExtractionError("No text content found in DOCX")

        return ExtractedDocument(
            title=source.filename or Path(file_path).stem,
            full_text=full_text,
            markdown="\n".join(markdown_lines),
            extraction_engine=self.engine_name,
            metadata={
                "source_type": "docx",
                "original_filename": source.filename or Path(file_path).name,
                "heading_count": heading_count,
                "paragraph_count": para_count,
            },
        )


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _paragraph_heading_level(style_name: str) -> int | None:
    """Return heading level (1-9) if *style_name* is a Word heading style, else None."""
    name_lower = style_name.lower().strip()
    # Match "Heading 1", "heading 2", "Heading 3 Title", etc.
    if name_lower.startswith("heading"):
        parts = name_lower.split()
        if len(parts) >= 2:
            try:
                level = int(parts[1])
                if 1 <= level <= 9:
                    return level
            except ValueError:
                pass
    return None
