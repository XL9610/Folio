"""PPTX extractor — preserves slide boundaries, titles, body text, and notes."""

from __future__ import annotations

from pathlib import Path

from .base import ContentExtractor, ExtractionError, ExtractorDependencyError
from .models import ExtractedDocument, SlideContent, SourceInput

_INSTALL_HINT = "pip install -r backend/requirements-extractors.txt"


class PptxExtractor(ContentExtractor):
    """Extract PPTX with slide boundaries, titles, body text, and speaker notes.

    SlideContent.slide_number is 1-based.
    Notes extraction failure is a warning, not total extraction failure.
    """

    engine_name = "python-pptx"

    # ------------------------------------------------------------------
    # ContentExtractor interface
    # ------------------------------------------------------------------

    def can_handle(self, source: SourceInput) -> bool:
        ext = Path(source.filename or source.file_path or "").suffix.lower()
        return source.source_type == "pptx" or ext == ".pptx"

    async def extract(self, source: SourceInput) -> ExtractedDocument:
        file_path = source.file_path
        if not file_path:
            raise ExtractionError("PPTX extraction requires a file_path")

        try:
            from pptx import Presentation  # type: ignore
        except ImportError:
            raise ExtractorDependencyError(
                f"python-pptx is required for PPTX extraction. Install with: {_INSTALL_HINT}"
            ) from None

        prs = Presentation(file_path)
        slides: list[SlideContent] = []
        warnings: list[str] = []
        all_text: list[str] = []

        for idx, slide in enumerate(prs.slides, start=1):
            title = _extract_slide_title(slide)
            body = _extract_slide_body(slide)
            notes = ""
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                warnings.append(f"Failed to extract notes from slide {idx}")

            sc = SlideContent(
                slide_number=idx,
                title=title,
                body_text=body,
                notes=notes,
            )
            slides.append(sc)
            slide_text = "\n".join(p for p in (title, body, notes) if p)
            if slide_text:
                all_text.append(slide_text)

        if not all_text:
            raise ExtractionError("No extractable text in PPTX slides")

        return ExtractedDocument(
            title=source.filename or Path(file_path).stem,
            full_text="\n".join(all_text),
            slides=slides,
            extraction_engine=self.engine_name,
            warnings=warnings,
            metadata={
                "source_type": "pptx",
                "original_filename": source.filename or Path(file_path).name,
                "slide_count": len(slides),
            },
        )


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _extract_slide_title(slide: object) -> str:
    """Extract title text from a slide's title placeholder."""
    try:
        if slide.shapes and slide.shapes.title:  # type: ignore
            return slide.shapes.title.text.strip()  # type: ignore
    except Exception:
        pass
    return ""


def _extract_slide_body(slide: object) -> str:
    """Extract visible text from all shapes on a slide."""
    parts: list[str] = []
    try:
        for shape in slide.shapes:  # type: ignore
            if hasattr(shape, "text") and shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    parts.append(text)
    except Exception:
        pass
    return "\n".join(parts)
